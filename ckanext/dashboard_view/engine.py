"""Bounded table preparation and dashboard queries, independent from CKAN.

Files are read only during preparation. Queries operate on a private immutable
DuckDB generation with external access disabled. User values are always bound;
identifiers must exist in the persisted column profile before they are quoted.
"""

import codecs
import csv
import datetime as dt
import io
import json
import math
import os
from pathlib import Path
import re
import tempfile
import threading
import zipfile

import duckdb

from .schema import normalize_config, normalize_source, validate_filters


DEFAULT_LIMITS = {"max_rows": 5_000_000, "max_columns": 500,
                  "max_bytes": 256 * 1024 * 1024, "memory_limit": "1GB", "threads": 2,
                  "max_xlsx_uncompressed_bytes": 1024 * 1024 * 1024}
PAGE_SIZE = 50
MAX_SCATTER_POINTS = 2000
MAX_MAP_POINTS = 2000
MAX_SERIES = 8
QUERY_TIMEOUT_SECONDS = 45
_ROW_KEY = "__dashboard_internal_row_id__"


def _quote(identifier):
    return '"%s"' % identifier.replace('"', '""')


def _json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _label(value, maximum=512):
    value = _json_value(value)
    if isinstance(value, str) and len(value) > maximum:
        return value[:maximum] + "…"
    return value


def _measure(value):
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("This calculation exceeds the supported numeric range. Narrow the filters or choose another aggregation.")
    return _json_value(value)


def _category_filter(widget, value):
    """Carry the real category identity separately from its display label."""
    if isinstance(value, str) and len(value) > 2000:
        return None
    if value is not None and widget["time_grain"] != "none":
        start = value.date() if isinstance(value, dt.datetime) else value
        if widget["time_grain"] == "day":
            end = start
        elif widget["time_grain"] == "month":
            import calendar
            end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
        else:
            end = start.replace(month=12, day=31)
        return {"field": widget["x_field"], "op": "between", "value": [start.isoformat(), end.isoformat()]}
    return {"field": widget["x_field"], "op": "eq", "value": _json_value(value)}


def _safe_rows(rows):
    return [[_json_value(value) for value in row] for row in rows]


def _limits(value):
    result = dict(DEFAULT_LIMITS)
    if value:
        result.update({key: val for key, val in value.items() if key in result})
    for key in ("max_rows", "max_columns", "max_bytes", "threads", "max_xlsx_uncompressed_bytes"):
        if isinstance(result[key], bool) or not isinstance(result[key], int) or result[key] < 1:
            raise ValueError("Invalid %s limit." % key)
    if not re.fullmatch(r"[1-9][0-9]*(?:MB|GB)", result["memory_limit"], re.IGNORECASE):
        raise ValueError("Invalid memory limit.")
    return result


def _column_names(header):
    used = {_ROW_KEY.casefold()}
    result = []
    for index, label in enumerate(header):
        label = str(label or "").strip().lstrip("\ufeff")
        if len(label) > 512 or "\x00" in label:
            raise ValueError("Column names must be at most 512 characters and contain no null bytes.")
        base = label or "column_%s" % (index + 1)
        name, suffix = base, 2
        while name.casefold() in used:
            name = "%s_%s" % (base, suffix)
            suffix += 1
        used.add(name.casefold())
        result.append({"key": name, "label": label or name})
    return result


def _detect_encoding(path, source):
    if source["encoding"]:
        return source["encoding"], []
    with open(path, "rb") as handle:
        sample = handle.read(131072)
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16", []
    try:
        codecs.getincrementaldecoder("utf-8-sig")().decode(sample, final=False)
        return "utf-8-sig", []
    except UnicodeDecodeError:
        return "cp1252", ["Windows-1252 encoding was detected. Confirm the encoding if accented characters look incorrect."]


def _csv_description(path, source, format_hint):
    encoding, warnings = _detect_encoding(path, source)
    try:
        with open(path, "r", encoding=encoding, newline="") as handle:
            sample = handle.read(131072)
            if not sample.strip():
                raise ValueError("The file is empty. Add a header row and data.")
            delimiter = source["delimiter"]
            if delimiter is None:
                if format_hint.lower() == "tsv":
                    delimiter = "\t"
                else:
                    # Score actual parsed rows. Sniffer's regex can mistake a
                    # semicolon inside a quoted header for the delimiter.
                    candidates = []
                    for candidate in (",", ";", "\t", "|"):
                        try:
                            parsed = list(csv.reader(io.StringIO(sample), delimiter=candidate))[:100]
                        except csv.Error:
                            continue
                        if parsed and len(parsed[0]) > 1:
                            body = parsed[1:-1] or parsed[1:] or parsed[:1]
                            consistency = sum(len(row) == len(parsed[0]) for row in body) / len(body)
                            candidates.append((consistency, len(parsed[0]), candidate))
                    delimiter = max(candidates)[2] if candidates else ","
            handle.seek(0)
            header = next(csv.reader(handle, delimiter=delimiter))
    except (UnicodeError, csv.Error) as exc:
        raise ValueError("The file cannot be read with this encoding and delimiter.") from exc
    return {"fields": _column_names(header), "encoding": encoding,
            "delimiter": delimiter, "warnings": warnings, "sheets": []}


def _check_xlsx(path, limits):
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 20000 or sum(member.file_size for member in members) > limits["max_xlsx_uncompressed_bytes"]:
                raise ValueError("The expanded spreadsheet exceeds the configured size limit.")
            if any(member.flag_bits & 1 for member in members):
                raise ValueError("Encrypted spreadsheets are not supported.")
            # read_only streams worksheets, but openpyxl still loads these
            # metadata collections into memory. Bound them independently.
            metadata_limits = {"xl/sharedStrings.xml": 64 * 1024 * 1024,
                               "xl/styles.xml": 8 * 1024 * 1024}
            if any(member.file_size > metadata_limits.get(member.filename, limits["max_xlsx_uncompressed_bytes"]) for member in members):
                raise ValueError("The spreadsheet contains too much shared text or style metadata. Export the selected sheet as CSV.")
    except zipfile.BadZipFile as exc:
        raise ValueError("The file is not a valid XLSX spreadsheet.") from exc


def inspect_file(input_path, source=None, format_hint="csv"):
    source = normalize_source(source)
    path = Path(input_path)
    if format_hint.lower().lstrip(".") == "xlsx" or path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        _check_xlsx(path, DEFAULT_LIMITS)
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            return {"sheets": workbook.sheetnames}
        finally:
            workbook.close()
    return _csv_description(path, source, format_hint)


def _materialize_xlsx(path, target, source, limits):
    from openpyxl import load_workbook
    _check_xlsx(path, limits)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("The XLSX spreadsheet cannot be opened.") from exc
    try:
        sheets = workbook.sheetnames
        sheet_name = source["sheet"] or sheets[0]
        if sheet_name not in sheets:
            raise ValueError("The selected worksheet is no longer available.")
        worksheet = workbook[sheet_name]
        if worksheet.max_column and worksheet.max_column > limits["max_columns"]:
            raise ValueError("The worksheet exceeds the column limit (%s)." % limits["max_columns"])
        if worksheet.max_row and worksheet.max_row > limits["max_rows"] + 1:
            raise ValueError("The worksheet exceeds the row limit (%s)." % limits["max_rows"])
        row_count = 0
        with open(target, "w", encoding="utf-8", newline="") as output:
            writer = csv.writer(output)
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index > limits["max_rows"]:
                    raise ValueError("The worksheet exceeds the row limit (%s)." % limits["max_rows"])
                if len(row) > limits["max_columns"]:
                    raise ValueError("The worksheet exceeds the column limit (%s)." % limits["max_columns"])
                writer.writerow([_json_value(value) if value is not None else "" for value in row])
                row_count += 1
                if output.tell() > limits["max_xlsx_uncompressed_bytes"]:
                    raise ValueError("The expanded spreadsheet exceeds the configured size limit.")
        if not row_count:
            raise ValueError("The selected worksheet is empty.")
        return sheets, sheet_name
    finally:
        workbook.close()


def _infer_type(values, source):
    values = [value.strip() for value in values if value is not None and str(value).strip()]
    if not values:
        return "text", None
    lower = {value.lower() for value in values}
    if lower <= {"true", "false", "yes", "no", "sí", "si"}:
        return "boolean", None
    # Leading zeros normally encode station/postal identifiers, not quantities.
    leading_zero = any(re.fullmatch(r"[+-]?0[0-9]+", value) for value in values)
    large_integer = any(re.fullmatch(r"[+-]?[0-9]+", value) and len(value.lstrip("+-0")) >= 16 and abs(int(value)) > 2 ** 53 for value in values if len(value) < 1000)
    numeric = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?")
    normalized = [value.replace(",", ".") if source["decimal"] == "," else value for value in values]
    if not leading_zero and not large_integer and all(numeric.fullmatch(value) and math.isfinite(float(value)) for value in normalized):
        return "number", None
    if source["date_format"]:
        try:
            for value in values:
                dt.datetime.strptime(value, source["date_format"])
            return "date", None
        except ValueError:
            pass
    if all(re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[T ])", value) for value in values):
        try:
            for value in values:
                dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
            return "date", None
        except ValueError:
            pass
    if any(re.fullmatch(r"\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}", value) for value in values):
        return "text", "Ambiguous dates were preserved as text. Set a date format to interpret this column."
    return "text", None


def _typed_expression(key, kind, source):
    column = _quote(key)
    trimmed = "nullif(trim(%s), '')" % column
    if kind == "text":
        return column, []
    if kind == "number":
        number = "replace(%s, ',', '.')" % trimmed if source["decimal"] == "," else trimmed
        cast = "try_cast(%s AS DOUBLE)" % number
        return "CASE WHEN isfinite(%s) THEN %s ELSE NULL END" % (cast, cast), []
    if kind == "boolean":
        return "CASE WHEN lower(%s) IN ('true','yes','sí','si','1') THEN TRUE WHEN lower(%s) IN ('false','no','0') THEN FALSE ELSE NULL END" % (trimmed, trimmed), []
    if source["date_format"]:
        return "try_cast(try_strptime(%s, ?) AS TIMESTAMPTZ) AT TIME ZONE 'UTC'" % trimmed, [source["date_format"]]
    return "try_cast(%s AS TIMESTAMPTZ) AT TIME ZONE 'UTC'" % trimmed, []


def prepare_table(input_path, output_path, source=None, format_hint="csv", limits=None):
    """Build original/typed row tables and return their JSON-serializable profile.

    A failed build deletes its partial output. An existing generation is never
    overwritten. The caller atomically publishes the completed database/profile.
    """
    source, limits = normalize_source(source), _limits(limits)
    input_path, output_path = Path(input_path), Path(output_path)
    if output_path.exists():
        raise ValueError("The output generation already exists.")
    if not input_path.is_file():
        raise ValueError("The source file is unavailable.")
    if input_path.stat().st_size > limits["max_bytes"]:
        raise ValueError("The source file exceeds the configured size limit.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_paths, connection = [], None
    warnings, sheets, selected_sheet = [], [], None
    try:
        csv_path = input_path
        csv_source = dict(source)
        if format_hint.lower().lstrip(".") == "xlsx" or input_path.suffix.lower() == ".xlsx":
            handle = tempfile.NamedTemporaryFile(prefix="dashboard-sheet-", suffix=".csv", dir=output_path.parent, delete=False)
            handle.close()
            csv_path = Path(handle.name)
            temporary_paths.append(csv_path)
            sheets, selected_sheet = _materialize_xlsx(input_path, csv_path, source, limits)
            csv_source.update({"encoding": "utf-8", "delimiter": ","})
            warnings.append("Spreadsheet formulas use their last saved values. Save the workbook in a spreadsheet application to refresh calculations.")
        description = _csv_description(csv_path, csv_source, format_hint)
        fields = description["fields"]
        if not fields or len(fields) > limits["max_columns"]:
            raise ValueError("The table must contain at least one column and remain within the column limit (%s)." % limits["max_columns"])
        warnings.extend(description["warnings"])
        if description["encoding"] not in {"utf-8", "utf-8-sig", "ascii"}:
            handle = tempfile.NamedTemporaryFile(prefix="dashboard-text-", suffix=".csv", dir=output_path.parent, delete=False)
            handle.close()
            normalized = Path(handle.name)
            temporary_paths.append(normalized)
            with open(csv_path, "r", encoding=description["encoding"], newline="") as original, open(normalized, "w", encoding="utf-8", newline="") as target:
                while True:
                    chunk = original.read(1024 * 1024)
                    if not chunk:
                        break
                    target.write(chunk)
            csv_path = normalized
        connection = duckdb.connect(str(output_path), config={"memory_limit": limits["memory_limit"], "threads": limits["threads"]})
        connection.execute("SET TimeZone='UTC'")
        connection.execute("CREATE TABLE original AS SELECT row_number() OVER () - 1 AS %s, * FROM read_csv(?, header=true, all_varchar=true, columns=?, delim=?, auto_detect=false, quote=?, escape=?, comment='', parallel=false, strict_mode=true, null_padding=false, max_line_size=4194304) LIMIT ?" % _quote(_ROW_KEY),
                           [str(csv_path), {field["key"]: "VARCHAR" for field in fields}, description["delimiter"], '"', '"', limits["max_rows"] + 1])
        rows = connection.execute("SELECT count(*) FROM original").fetchone()[0]
        if rows > limits["max_rows"]:
            raise ValueError("The table exceeds the row limit (%s)." % limits["max_rows"])
        examples = connection.execute("SELECT %s FROM original LIMIT 1000" % ",".join(_quote(field["key"]) for field in fields)).fetchall()
        expressions, expression_parameters = [_quote(_ROW_KEY)], []
        field_keys = {field["key"] for field in fields}
        for unknown in source["types"].keys() - field_keys:
            warnings.append("Type override column '%s' is no longer available." % unknown)
        for index, field in enumerate(fields):
            values = [row[index] for row in examples]
            inferred, warning = _infer_type(values, source)
            kind = source["types"].get(field["key"], inferred)
            if warning and field["key"] not in source["types"]:
                warnings.append("%s: %s" % (field["label"], warning))
            field.update({"type": kind, "examples": list(dict.fromkeys(value[:200] for value in values if value is not None))[:5]})
            expression, parameters = _typed_expression(field["key"], kind, source)
            expressions.append("%s AS %s" % (expression, _quote(field["key"])))
            expression_parameters.extend(parameters)
        connection.execute("CREATE TABLE data AS SELECT %s FROM original" % ",".join(expressions), expression_parameters)
        # Measure conversion loss over the complete table, not only the sample.
        typed_fields = [field for field in fields if field["type"] != "text"]
        if typed_fields:
            checks = ["count(*) FILTER (WHERE nullif(trim(o.%s), '') IS NOT NULL AND d.%s IS NULL)" % (_quote(field["key"]), _quote(field["key"])) for field in typed_fields]
            losses = connection.execute("SELECT %s FROM original o JOIN data d USING (%s)" % (",".join(checks), _quote(_ROW_KEY))).fetchone()
            for field, loss in zip(typed_fields, losses):
                field["invalid_count"] = loss
                if loss:
                    warnings.append("%s: %s values could not be interpreted as %s; they are excluded from calculations and preserved in CSV export." % (field["label"], loss, field["type"]))
        profile = {"rows": rows, "fields": fields, "sheets": sheets, "sheet": selected_sheet,
                   "warnings": warnings, "encoding": description["encoding"], "delimiter": description["delimiter"],
                   "source": source, "updated_at": dt.datetime.now(dt.timezone.utc).isoformat()}
        connection.execute("CREATE TABLE __dashboard_profile (profile VARCHAR)")
        connection.execute("INSERT INTO __dashboard_profile VALUES (?)", [json.dumps(profile, allow_nan=False)])
        connection.execute("CHECKPOINT")
        connection.close()
        connection = None
        os.chmod(output_path, 0o600)
        return profile
    except Exception as exc:
        if connection is not None:
            connection.close()
        for path in (output_path, Path(str(output_path) + ".wal")):
            path.unlink(missing_ok=True)
        if isinstance(exc, ValueError):
            raise
        if isinstance(exc, UnicodeError):
            raise ValueError("The source contains characters incompatible with the selected encoding.") from exc
        if isinstance(exc, duckdb.Error):
            raise ValueError("The table could not be prepared. Check the delimiter, consistent row lengths, and configured memory limits.") from exc
        raise ValueError("The source could not be read as a supported table.") from exc
    finally:
        for path in temporary_paths:
            path.unlink(missing_ok=True)


def _open_generation(path):
    connection = duckdb.connect(str(path), read_only=True,
                                config={"enable_external_access": "false", "memory_limit": "768MB", "threads": 2})
    try:
        connection.execute("SET TimeZone='UTC'")
        profile = json.loads(connection.execute("SELECT profile FROM __dashboard_profile LIMIT 1").fetchone()[0])
        return connection, profile
    except Exception:
        connection.close()
        raise


def _filter_value(value, kind):
    if value is None:
        return None
    if kind == "number":
        try:
            converted = float(value)
            if not math.isfinite(converted):
                raise ValueError()
            return converted
        except (ValueError, TypeError) as exc:
            raise ValueError("Use a valid number for this filter.") from exc
    if kind == "date":
        try:
            converted = dt.datetime.fromisoformat(str(value).replace("Z", "+00:00"))
            if converted.tzinfo is not None:
                converted = converted.astimezone(dt.timezone.utc)
            return converted.replace(tzinfo=None)
        except ValueError as exc:
            raise ValueError("Use an ISO date or datetime for this filter.") from exc
    if kind == "boolean":
        if value is True or str(value).lower() in {"true", "1", "yes", "sí", "si"}:
            return True
        if value is False or str(value).lower() in {"false", "0", "no"}:
            return False
        raise ValueError("Use true or false for this filter.")
    return str(value)


def _where(filters, fields, config, bounds=None, prefix=""):
    expressions, parameters = [], []
    for item in validate_filters(filters, fields):
        field, op = item["field"], item["op"]
        column = prefix + _quote(field)
        values = item["value"] if op in {"in", "between"} else [item["value"]]
        date_only_upper = fields[field]["type"] == "date" and op in {"lte", "between"} and isinstance(values[-1], str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", values[-1])
        values = [_filter_value(value, fields[field]["type"]) for value in values]
        if op == "in":
            non_null = [value for value in values if value is not None]
            parts = []
            if non_null:
                parts.append("%s IN (%s)" % (column, ",".join("?" for _ in non_null)))
                parameters.extend(non_null)
            if None in values:
                parts.append("%s IS NULL" % column)
            expressions.append("(" + " OR ".join(parts) + ")" if parts else "FALSE")
        elif op == "eq" and values[0] is None:
            expressions.append("%s IS NULL" % column)
        elif op == "between":
            if date_only_upper:
                expressions.append("(%s >= ? AND %s < ?)" % (column, column))
                values[-1] += dt.timedelta(days=1)
            else:
                expressions.append("%s BETWEEN ? AND ?" % column)
            parameters.extend(values)
        elif op == "lte" and date_only_upper:
            expressions.append("%s < ?" % column)
            parameters.append(values[0] + dt.timedelta(days=1))
        else:
            expressions.append("%s %s ?" % (column, {"eq": "=", "gte": ">=", "lte": "<="}[op]))
            parameters.extend(values)
    if bounds is not None:
        if not isinstance(bounds, dict):
            raise ValueError("Map bounds must be an object.")
        try:
            west, south, east, north = [float(bounds[name]) for name in ("west", "south", "east", "north")]
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError("Invalid map bounds.") from exc
        if not all(math.isfinite(value) for value in (west, south, east, north)) or not (-180 <= west <= 180 and -180 <= east <= 180 and -90 <= south <= north <= 90):
            raise ValueError("Invalid map bounds.")
        widget_id = bounds.get("widget_id")
        if widget_id is not None and (not isinstance(widget_id, str) or len(widget_id) > 80):
            raise ValueError("Invalid map block identifier.")
        map_widget = next((widget for widget in config["widgets"] if widget["type"] == "map" and (widget_id is None or widget["id"] == widget_id) and widget.get("lat_field") in fields and widget.get("lon_field") in fields and fields[widget["lat_field"]]["type"] == fields[widget["lon_field"]]["type"] == "number"), None)
        if not map_widget:
            raise ValueError("A map with valid coordinate columns is required to filter an area.")
        lat, lon = [prefix + _quote(map_widget[name]) for name in ("lat_field", "lon_field")]
        expressions.append("%s BETWEEN ? AND ?" % lat)
        parameters.extend([south, north])
        expressions.append("(%s >= ? OR %s <= ?)" % (lon, lon) if west > east else "%s BETWEEN ? AND ?" % lon)
        parameters.extend([west, east])
    return (" AND ".join(expressions) or "TRUE"), parameters


def _require(widget, name, fields, types=None):
    key = widget.get(name)
    if not key:
        raise ValueError("Choose a %s column." % name.replace("_field", "").replace("_", " "))
    if key not in fields:
        raise ValueError("Column '%s' is no longer available." % key)
    if types and fields[key]["type"] not in types:
        raise ValueError("Column '%s' has an incompatible type. Update this block or its column interpretation." % key)
    return _quote(key)


def _aggregate(widget, fields, prefix=""):
    aggregate = widget["aggregate"]
    if aggregate == "count":
        return "count(*)"
    types = None if aggregate == "count_distinct" else {"number"}
    if widget["type"] == "kpi" and aggregate in {"min", "max"}:
        types = {"number", "date"}
    column = prefix + _require(widget, "y_field", fields, types)
    if aggregate == "count_distinct":
        return "count(DISTINCT %s)" % column
    return "%s(%s)" % (aggregate, column)


def _table_result(connection, widget, fields, where, params, count, pages):
    columns = widget["columns"] or list(fields)
    if any(column not in fields for column in columns):
        raise ValueError("A table column is no longer available. Update the selected columns.")
    page = (pages or {}).get(widget["id"], 1)
    if isinstance(page, bool) or not isinstance(page, int) or page < 1 or page > 1000000:
        raise ValueError("Invalid table page.")
    page = min(page, max(1, math.ceil(count / PAGE_SIZE)))
    order = _quote(_ROW_KEY)
    if widget["x_field"]:
        order = _require(widget, "x_field", fields) + (" ASC" if widget["sort"] == "asc" else " DESC") + " NULLS LAST, " + order
    selected = ["left(%s, 2001)" % _quote(column) if fields[column]["type"] == "text" else _quote(column) for column in columns]
    rows = connection.execute("SELECT %s FROM data WHERE %s ORDER BY %s LIMIT ? OFFSET ?" % (",".join(selected), where, order), params + [PAGE_SIZE, (page - 1) * PAGE_SIZE]).fetchall()
    truncated = any(isinstance(value, str) and len(value) > 2000 for row in rows for value in row)
    rows = [[value[:2000] + "…" if isinstance(value, str) and len(value) > 2000 else value for value in row] for row in rows]
    return {"type": "table", "columns": [{"key": key, "label": fields[key]["label"]} for key in columns],
            "rows": [dict(zip(columns, row)) for row in _safe_rows(rows)], "total": count, "page": page, "page_size": PAGE_SIZE, "truncated": truncated}


def _chart_result(connection, widget, fields, where, params, count):
    x = _require(widget, "x_field", fields)
    if widget["time_grain"] != "none":
        _require(widget, "x_field", fields, {"date"})
        x = "date_trunc('%s', %s)" % (widget["time_grain"], x)
    aggregate = _aggregate(widget, fields)
    sort = "ASC" if widget["sort"] == "asc" else "DESC"
    order = "x ASC NULLS LAST" if widget["type"] in {"line", "area"} else "value %s NULLS LAST, x ASC NULLS LAST" % sort
    groups = connection.execute("SELECT %s AS x, %s AS value FROM data WHERE %s GROUP BY 1 ORDER BY %s LIMIT ?" % (x, aggregate, where, order), params + [widget["limit"] + 1]).fetchall()
    truncated = len(groups) > widget["limit"]
    groups = groups[:widget["limit"]]
    labels = [_label(row[0]) if row[0] is not None else "(Empty)" for row in groups]
    text_truncated = any(isinstance(row[0], str) and len(row[0]) > 512 for row in groups)
    if not widget["series_field"]:
        datasets = [{"label": widget["title"] or widget["aggregate"], "data": [_measure(row[1]) for row in groups]}]
    elif not groups:
        _require(widget, "series_field", fields)
        datasets = []
    else:
        series = _require(widget, "series_field", fields)
        values = [row[0] for row in groups]
        dimension_conditions = ["%s IS NOT DISTINCT FROM ?" % x for _ in values]
        narrowed = "(%s) AND (%s)" % (where, " OR ".join(dimension_conditions))
        narrowed_params = params + values
        series_rows = connection.execute("SELECT %s AS series, count(*) AS n FROM data WHERE %s GROUP BY 1 ORDER BY n DESC, series ASC NULLS LAST LIMIT ?" % (series, narrowed), narrowed_params + [MAX_SERIES + 1]).fetchall()
        truncated = truncated or len(series_rows) > MAX_SERIES
        series_values = [row[0] for row in series_rows[:MAX_SERIES]]
        series_conditions = " OR ".join("%s IS NOT DISTINCT FROM ?" % series for _ in series_values)
        records = connection.execute("SELECT %s AS x, %s AS series, %s AS value FROM data WHERE %s AND (%s) GROUP BY 1, 2" % (x, series, aggregate, narrowed, series_conditions), narrowed_params + series_values).fetchall()
        lookup = {(record[0], record[1]): _measure(record[2]) for record in records}
        missing = 0 if widget["aggregate"] in {"count", "count_distinct"} else None
        text_truncated = text_truncated or any(isinstance(value, str) and len(value) > 512 for value in series_values)
        datasets = [{"label": str(_label(value)) if value is not None else "(Empty)",
                     "data": [lookup.get((group[0], value), missing) for group in groups]} for value in series_values]
    return {"type": widget["type"], "labels": labels, "datasets": datasets, "count": count, "truncated": truncated,
            "category_filters": [_category_filter(widget, row[0]) for row in groups],
            "text_truncated": text_truncated}


def _histogram_result(connection, widget, fields, where, params, count):
    field_name = "x_field" if widget["x_field"] else "y_field"
    column = _require(widget, field_name, fields, {"number"})
    low, high, valid_count = connection.execute("SELECT min(%s), max(%s), count(%s) FROM data WHERE %s" % (column, column, column, where), params).fetchone()
    if not valid_count:
        return {"type": "histogram", "labels": [], "datasets": [], "count": count, "valid_count": 0}
    bins = widget["bins"] if high != low else 1
    width = (high - low) / bins if bins > 1 else 1
    if width == 0:
        bins, width = 1, 1
    if math.isfinite(width):
        expression = "least(?, greatest(0, floor((%s - ?) / ?)::INTEGER))" % column
        bucket_parameters = [bins - 1, low, width]
    else:
        expression = "least(?, greatest(0, floor(((%s / 2 - ?) / ?) * ?)::INTEGER))" % column
        bucket_parameters = [bins - 1, low / 2, high / 2 - low / 2, bins]
    bucket_rows = connection.execute("SELECT %s AS bucket, count(*) FROM data WHERE (%s) AND %s IS NOT NULL GROUP BY 1 ORDER BY bucket" % (expression, where, column), bucket_parameters + params).fetchall()
    bucket_counts = dict(bucket_rows)
    edges = [(1 - index / bins) * low + (index / bins) * high for index in range(bins + 1)]
    labels = ["%s – %s" % (format(edges[index], ".6g"), format(edges[index + 1], ".6g")) for index in range(bins)]
    if high == low:
        labels = [format(low, ".6g")]
    return {"type": "histogram", "labels": labels, "datasets": [{"label": widget["title"] or "Count", "data": [bucket_counts.get(index, 0) for index in range(bins)]}], "count": count, "valid_count": valid_count}


def _scatter_result(connection, widget, fields, where, params, count):
    x, y = [_require(widget, key, fields, {"number"}) for key in ("x_field", "y_field")]
    series = _require(widget, "series_field", fields) if widget["series_field"] else "NULL"
    valid_where = "(%s) AND %s IS NOT NULL AND %s IS NOT NULL" % (where, x, y)
    valid_count = connection.execute("SELECT count(*) FROM data WHERE " + valid_where, params).fetchone()[0]
    base = "SELECT %s AS x, %s AS y, %s AS series FROM data WHERE %s" % (x, y, series, valid_where)
    sampled = valid_count > MAX_SCATTER_POINTS
    query = "SELECT * FROM (%s) USING SAMPLE reservoir(%s ROWS) REPEATABLE(42)" % (base, MAX_SCATTER_POINTS) if sampled else base
    rows = connection.execute(query, params).fetchall()
    datasets = {}
    # Keep all points, folding rare series into Other with explicit disclosure.
    series_counts = {}
    for _, _, label in rows:
        series_counts[label] = series_counts.get(label, 0) + 1
    selected = set(sorted(series_counts, key=lambda value: -series_counts[value])[:MAX_SERIES])
    truncated = len(series_counts) > MAX_SERIES
    for x_value, y_value, label in rows:
        label = "Other series" if label not in selected else str(_label(label)) if label is not None else widget["title"] or "Observations"
        datasets.setdefault(label, []).append({"x": x_value, "y": y_value})
    return {"type": "scatter", "datasets": [{"label": key, "data": values} for key, values in datasets.items()], "count": count, "valid_count": valid_count, "sampled": sampled, "truncated": truncated}


def _map_result(connection, widget, fields, where, params, count, bounds):
    lat, lon = [_require(widget, key, fields, {"number"}) for key in ("lat_field", "lon_field")]
    valid_where = "(%s) AND %s BETWEEN -90 AND 90 AND %s BETWEEN -180 AND 180" % (where, lat, lon)
    valid_count = connection.execute("SELECT count(*) FROM data WHERE " + valid_where, params).fetchone()[0]
    if not valid_count:
        return {"type": "map", "points": [], "count": count, "valid_count": 0, "aggregated": False}
    if valid_count <= MAX_MAP_POINTS:
        if widget["aggregate"] == "count":
            value = "1"
        elif widget["aggregate"] == "count_distinct":
            value = "CASE WHEN %s IS NULL THEN 0 ELSE 1 END" % _require(widget, "y_field", fields)
        else:
            value = _require(widget, "y_field", fields, {"number"})
        rows = connection.execute("SELECT %s, %s, %s FROM data WHERE %s ORDER BY %s" % (lat, lon, value, valid_where, _quote(_ROW_KEY)), params).fetchall()
        points = [{"lat": row[0], "lon": row[1], "value": _measure(row[2]), "count": 1} for row in rows]
        return {"type": "map", "points": points, "count": count, "valid_count": valid_count, "aggregated": False}
    # Cell size derives from the actual extent, bounding cells to <= 46 * 23.
    south, north, west, east = connection.execute("SELECT min(%s),max(%s),min(%s),max(%s) FROM data WHERE %s" % (lat, lat, lon, lon, valid_where), params).fetchone()
    lat_size = max((north - south) / 22, 0.000001)
    lon_size = max((east - west) / 44, 0.000001)
    aggregate = _aggregate(widget, fields)
    rows = connection.execute("SELECT avg(%s), avg(%s), %s, count(*) FROM data WHERE %s GROUP BY floor((%s - ?) / ?), floor((%s - ?) / ?)" % (lat, lon, aggregate, valid_where, lat, lon), params + [south, lat_size, west, lon_size]).fetchall()
    points = [{"lat": row[0], "lon": row[1], "value": _measure(row[2]), "count": row[3]} for row in rows]
    return {"type": "map", "points": points, "count": count, "valid_count": valid_count, "aggregated": True,
            "cell_size": {"latitude": lat_size, "longitude": lon_size}}


def _facets(connection, config, fields, filters, bounds):
    facets, ranges, truncated = {}, {}, []
    for definition in config["filters"]:
        key = definition["field"]
        if key not in fields:
            continue
        column = _quote(key)
        where, params = _where([item for item in (filters or []) if item["field"] != key], fields, config, bounds)
        if definition["type"] == "category":
            if fields[key]["type"] == "text":
                long_values = connection.execute("SELECT EXISTS(SELECT 1 FROM data WHERE (%s) AND length(%s) > 2000)" % (where, column), params).fetchone()[0]
                if long_values:
                    truncated.append(key)
                where = "(%s) AND (%s IS NULL OR length(%s) <= 2000)" % (where, column, column)
            rows = connection.execute("SELECT %s AS value, count(*) AS n FROM data WHERE %s GROUP BY 1 ORDER BY n DESC, value ASC NULLS LAST LIMIT 101" % (column, where), params).fetchall()
            if len(rows) > 100 and key not in truncated:
                truncated.append(key)
            facets[key] = [{"value": _json_value(value), "count": count} for value, count in rows[:100]]
        elif fields[key]["type"] in {"number", "date"}:
            low, high = connection.execute("SELECT min(%s),max(%s) FROM data WHERE %s" % (column, column, where), params).fetchone()
            ranges[key] = {"min": _json_value(low), "max": _json_value(high)}
    return {"facets": facets, "ranges": ranges, "facets_truncated": truncated}


def query_table(db_path, config, filters=None, bounds=None, pages=None):
    config = normalize_config(config, allow_empty=True)
    if pages is not None and (not isinstance(pages, dict) or len(pages) > 24):
        raise ValueError("Table pages must be an object.")
    connection, profile = _open_generation(db_path)
    timer = threading.Timer(QUERY_TIMEOUT_SECONDS, connection.interrupt)
    timer.daemon = True
    timer.start()
    try:
        fields = {field["key"]: field for field in profile["fields"]}
        where, params = _where(filters, fields, config, bounds)
        count = connection.execute("SELECT count(*) FROM data WHERE " + where, params).fetchone()[0]
        results = {}
        for widget in config["widgets"]:
            try:
                kind = widget["type"]
                if kind == "text":
                    result = {"type": "text", "text": widget["text"]}
                elif kind == "kpi":
                    value = connection.execute("SELECT %s FROM data WHERE %s" % (_aggregate(widget, fields), where), params).fetchone()[0]
                    result = {"type": "kpi", "value": _measure(value), "count": count}
                elif kind == "table":
                    result = _table_result(connection, widget, fields, where, params, count, pages)
                elif kind == "histogram":
                    result = _histogram_result(connection, widget, fields, where, params, count)
                elif kind == "scatter":
                    result = _scatter_result(connection, widget, fields, where, params, count)
                elif kind == "map":
                    result = _map_result(connection, widget, fields, where, params, count, bounds)
                else:
                    result = _chart_result(connection, widget, fields, where, params, count)
            except ValueError as exc:
                result = {"type": widget["type"], "error": str(exc), "code": "invalid_widget"}
            results[widget["id"]] = result
        response = {"rows": count, "results": results, "warnings": profile.get("warnings", [])}
        response.update(_facets(connection, config, fields, filters, bounds))
        return response
    except duckdb.InterruptException as exc:
        raise ValueError("The dashboard query exceeded its time limit. Reduce the number of blocks or narrow the filters.") from exc
    except duckdb.Error as exc:
        raise ValueError("The dashboard query could not complete within its data and memory limits.") from exc
    finally:
        timer.cancel()
        timer.join()
        connection.close()


def _export_value(value):
    """Preserve original content while preventing spreadsheet formula execution."""
    if value is None:
        return ""
    value = str(value)
    stripped = value.lstrip(" \t\r\n")
    if stripped.startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        # Negative numeric originals remain useful as numbers; formula-like
        # strings, including whitespace-prefixed variants, become literal text.
        if not re.fullmatch(r"-[0-9]+(?:[.,][0-9]+)?(?:[eE][+-]?[0-9]+)?", stripped):
            return "'" + value
    return value


def export_csv(db_path, config, filters=None, bounds=None):
    """Yield bounded UTF-8 chunks of all original filtered columns.

    The caller must authorize before consuming the iterator. The connection is
    closed when streaming finishes or the client closes the generator.
    """
    config = normalize_config(config, allow_empty=True)
    connection, profile = _open_generation(db_path)
    timer = threading.Timer(120, connection.interrupt)
    timer.daemon = True
    timer.start()
    try:
        fields = {field["key"]: field for field in profile["fields"]}
        where, params = _where(filters, fields, config, bounds, prefix="d.")
        selected = ",".join("o." + _quote(key) for key in fields)
        connection.execute("SELECT %s FROM original o JOIN data d USING (%s) WHERE %s ORDER BY o.%s" % (selected, _quote(_ROW_KEY), where, _quote(_ROW_KEY)), params)
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([_export_value(field["label"]) for field in fields.values()])
        yield output.getvalue().encode("utf-8")
        output.seek(0)
        output.truncate(0)
        while True:
            rows = connection.fetchmany(250)
            if not rows:
                break
            for row in rows:
                writer.writerow([_export_value(value) for value in row])
                if output.tell() >= 65536:
                    yield output.getvalue().encode("utf-8")
                    output.seek(0)
                    output.truncate(0)
            if output.tell():
                yield output.getvalue().encode("utf-8")
                output.seek(0)
                output.truncate(0)
    finally:
        timer.cancel()
        timer.join()
        connection.close()
